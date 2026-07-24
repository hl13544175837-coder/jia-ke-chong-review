#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""将 PDF / DOCX 文档逐页转换为图片 data URI，统一送入视觉模型管线。"""

from __future__ import annotations

import base64
import io
import math
from pathlib import Path
from typing import List, Tuple

import fitz  # PyMuPDF
from docx import Document
from PIL import Image, ImageDraw, ImageFont

# ── 文档 → 图片 常量 ─────────────────────────────────────────────
PAGE_WIDTH = 1190
PAGE_HEIGHT = 1684  # A4 比例
PAGE_MARGIN = 60
FONT_SIZE = 20
LINE_HEIGHT = 30
FONT_COLOR = (0, 0, 0)
BG_COLOR = (255, 255, 255)

# ── PDF ──────────────────────────────────────────────────────────


def pdf_to_images(file_path: Path) -> List[Tuple[bytes, str]]:
    """将 PDF 每页渲染为 JPEG 字节，返回 [(bytes, "image/jpeg")]."""
    pages: list[tuple[bytes, str]] = []
    doc = fitz.open(str(file_path))
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        pages.append((buf.getvalue(), "image/jpeg"))
    doc.close()
    return pages


# ── DOCX ─────────────────────────────────────────────────────────


def docx_to_images(file_path: Path) -> List[Tuple[bytes, str]]:
    """将 DOCX 文本内容渲染为图片页，返回 [(bytes, "image/jpeg")]."""
    doc = Document(str(file_path))
    lines = _extract_docx_lines(doc)

    font = _load_font()
    pages = _layout_pages(lines, font)
    return [(page_bytes(img), "image/jpeg") for img in pages]


def _extract_docx_lines(doc: Document) -> list[str]:
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return lines


def _load_font() -> ImageFont.FreeTypeFont | None:
    # 优先用系统自带中文字体
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",       # 微软雅黑
        "C:/Windows/Fonts/simsun.ttc",      # 宋体
        "C:/Windows/Fonts/simhei.ttf",      # 黑体
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, FONT_SIZE)
    return None


def _layout_pages(lines: list[str], font: ImageFont.FreeTypeFont | None) -> list[Image.Image]:
    draw_area_width = PAGE_WIDTH - 2 * PAGE_MARGIN
    draw_area_height = PAGE_HEIGHT - 2 * PAGE_MARGIN
    pages: list[Image.Image] = []
    y = PAGE_MARGIN
    img = _new_page()

    for line in lines:
        wrapped = _wrap_line(line, font, draw_area_width)
        for wline in wrapped:
            if y + LINE_HEIGHT > PAGE_HEIGHT - PAGE_MARGIN:
                pages.append(img)
                img = _new_page()
                y = PAGE_MARGIN
            draw = ImageDraw.Draw(img)
            draw.text((PAGE_MARGIN, y), wline, font=font, fill=FONT_COLOR)
            y += LINE_HEIGHT

    pages.append(img)
    return pages


def _new_page() -> Image.Image:
    return Image.new("RGB", (PAGE_WIDTH, PAGE_HEIGHT), BG_COLOR)


def _wrap_line(text: str, font: ImageFont.FreeTypeFont | None, max_width: int) -> list[str]:
    """简单按字符折行，PIL 不支持中英文混排自动断行。"""
    if not text:
        return [""]
    chars_per_line = max_width // FONT_SIZE
    lines: list[str] = []
    for i in range(0, len(text), chars_per_line):
        lines.append(text[i : i + chars_per_line])
    return lines


def page_bytes(img: Image.Image, quality: int = 92) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=quality)
    return buf.getvalue()


# ── 统一入口 ─────────────────────────────────────────────────────


def document_to_image_data_uris(file_path: Path) -> List[str]:
    """将 PDF 或 DOCX 逐页转为 data URI 列表，可直接送给视觉模型。"""
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        pages = pdf_to_images(file_path)
    elif ext == ".docx":
        pages = docx_to_images(file_path)
    else:
        raise ValueError(f"不支持的文档格式: {ext}")

    uris: list[str] = []
    for img_bytes, mime in pages:
        encoded = base64.b64encode(img_bytes).decode("ascii")
        uris.append(f"data:{mime};base64,{encoded}")
    return uris
